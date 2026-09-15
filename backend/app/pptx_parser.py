"""Extract everything gradeable from a narrated PPTX.

A narrated PPTX (PowerPoint "Record Slide Show") stores one audio file per
slide under ppt/media/, linked from the slide's relationship part. We pull out,
per slide: visible text, speaker notes, the audio file (if any) with its
duration, and optionally a rendered PNG of the slide.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field, asdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

AUDIO_EXT = {".mp3", ".m4a", ".wav", ".wma", ".aac", ".mp4", ".m4v", ".ogg", ".oga", ".opus", ".flac", ".webm", ".mov", ".avi", ".wmv", ".mid"}


@dataclass
class SlideData:
    number: int
    title: str
    text: str
    notes: str
    audio_path: Path | None = None
    audio_duration_s: float | None = None
    image_path: Path | None = None
    transcript: str | None = None
    word_count: int = 0  # of the transcript, filled after transcription

    def to_dict(self) -> dict:
        d = asdict(self)
        for k in ("audio_path", "image_path"):
            d[k] = str(d[k]) if d[k] else None
        return d


@dataclass
class Deck:
    source: Path
    slides: list[SlideData] = field(default_factory=list)
    work_dir: Path | None = None

    @property
    def total_audio_s(self) -> float:
        return sum(s.audio_duration_s or 0.0 for s in self.slides)

    @property
    def narrated_slides(self) -> int:
        return sum(1 for s in self.slides if s.audio_path)

    def to_dict(self) -> dict:
        return {
            "source": self.source.name,
            "slide_count": len(self.slides),
            "narrated_slides": self.narrated_slides,
            "total_audio_s": round(self.total_audio_s, 1),
            "slides": [s.to_dict() for s in self.slides],
        }


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
        for sub in shape.shapes:
            out.extend(_shape_text(sub))
    return out


def _slide_title(slide) -> str:
    try:
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    # fall back to the top-most text shape
    best = None
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            top = shape.top if shape.top is not None else Emu(10**12)
            if best is None or top < best[0]:
                best = (top, shape.text_frame.text.strip().splitlines()[0])
    return best[1] if best else ""


def _slide_audio(slide, work_dir: Path, number: int) -> Path | None:
    """Return the first audio media part linked from this slide, copied to work_dir."""
    for rel in slide.part.rels.values():
        if rel.is_external:
            continue
        target = rel.target_ref
        ext = Path(target).suffix.lower()
        if ext in AUDIO_EXT and ("media" in rel.reltype or "audio" in rel.reltype or "video" in rel.reltype):
            blob = rel.target_part.blob
            out = work_dir / f"slide{number:02d}{ext}"
            out.write_bytes(blob)
            return out
    return None


def to_mp3(path: Path) -> Path:
    """Re-encode any embedded audio/video track to a mono 16 kHz MP3 (what the STT model accepts).
    Returns the original path if ffmpeg is unavailable or fails."""
    if not shutil.which("ffmpeg"):
        return path
    out = path.with_name(path.stem + "_stt.mp3")
    try:
        subprocess.run(
            ["ffmpeg", "-y", "-v", "error", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", "-b:a", "48k", str(out)],
            capture_output=True, check=True, timeout=300,
        )
        return out if out.exists() and out.stat().st_size > 0 else path
    except Exception:
        return path


def audio_duration(path: Path) -> float | None:
    if not shutil.which("ffprobe"):
        return None
    try:
        r = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)],
            capture_output=True, text=True, check=True, timeout=60,
        )
        return round(float(json.loads(r.stdout)["format"]["duration"]), 2)
    except Exception:
        return None


def render_slides(pptx: Path, out_dir: Path, dpi: int = 72) -> list[Path]:
    """PPTX -> PDF (LibreOffice) -> PNG per page (pdftoppm). Returns [] if tools are missing."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        return []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(pptx)],
            capture_output=True, check=True, timeout=300,
        )
        pdfs = list(Path(tmp).glob("*.pdf"))
        if not pdfs:
            return []
        subprocess.run(
            [pdftoppm, "-png", "-r", str(dpi), str(pdfs[0]), str(out_dir / "slide")],
            capture_output=True, check=True, timeout=300,
        )
    return sorted(out_dir.glob("slide-*.png"))


def parse_pptx(pptx: Path, work_dir: Path, render: bool = True) -> Deck:
    work_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(pptx))
    deck = Deck(source=pptx, work_dir=work_dir)

    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        audio = _slide_audio(slide, work_dir, i)
        if audio:
            audio = to_mp3(audio)
        deck.slides.append(
            SlideData(
                number=i,
                title=_slide_title(slide),
                text="\n".join(lines),
                notes=notes,
                audio_path=audio,
                audio_duration_s=audio_duration(audio) if audio else None,
            )
        )

    if render:
        images = render_slides(pptx, work_dir)
        for s, img in zip(deck.slides, images):
            s.image_path = img

    return deck
