"""Render a GradingResult as Markdown and DOCX, in German or English."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .grader import FormalCheck, GradingResult
from .pptx_parser import Deck

T = {
    "de": {
        "title": "Bewertung", "total": "Gesamt", "points": "Punkte", "recommendation": "Empfehlung der Gruppe",
        "criteria": "Kriterien", "criterion": "Kriterium", "level": "Niveau",
        "holds": "Was trägt.", "thin": "Was bleibt dünn.", "next": "Nächster Schritt.", "evidence": "Belege",
        "strengths": "Stärken", "weaknesses": "Schwächen", "coverage": "Abdeckung der Aufgaben",
        "task": "Aufgabe", "status": "Status", "note": "Anmerkung", "delivery": "Vortrag (Tonspur)",
        "formal": "Formale Checks", "metrics": "Kennzahlen", "metric": "Kennzahl", "value": "Wert",
        "slides": "Folien (vertont)", "duration": "Dauer Tonspur", "wpm": "Sprechtempo", "wpm_unit": "Wörter/min",
        "fillers": "Füllwörter", "slide_notes": "Anmerkungen je Folie", "slide": "Folie", "flags": "Manuell prüfen",
        "levels": {"ueberzeugend": "überzeugend", "tragfaehig": "tragfähig", "ansatzweise": "ansatzweise"},
        "statuses": {"adressiert": "adressiert", "teilweise": "teilweise", "fehlt": "fehlt"},
    },
    "en": {
        "title": "Assessment", "total": "Total", "points": "points", "recommendation": "The group's recommendation",
        "criteria": "Criteria", "criterion": "Criterion", "level": "Level",
        "holds": "What holds.", "thin": "What stays thin.", "next": "Next step.", "evidence": "Evidence",
        "strengths": "Strengths", "weaknesses": "Weaknesses", "coverage": "Coverage of the assignments",
        "task": "Assignment", "status": "Status", "note": "Note", "delivery": "Delivery (audio track)",
        "formal": "Formal checks", "metrics": "Metrics", "metric": "Metric", "value": "Value",
        "slides": "Slides (narrated)", "duration": "Audio duration", "wpm": "Speaking rate", "wpm_unit": "words/min",
        "fillers": "Filler words", "slide_notes": "Notes per slide", "slide": "Slide", "flags": "Check manually",
        "levels": {"ueberzeugend": "convincing", "tragfaehig": "sound", "ansatzweise": "rudimentary"},
        "statuses": {"adressiert": "addressed", "teilweise": "partly", "fehlt": "missing"},
    },
}
OK = {True: "✓", False: "✗", None: "–"}


def _metrics_rows(deck: Deck, metrics: dict, t: dict) -> list[list[str]]:
    return [
        [t["slides"], f"{len(deck.slides)} ({deck.narrated_slides})"],
        [t["duration"], f"{metrics['total_duration_min']} min"],
        [t["wpm"], f"{metrics['words_per_minute']} {t['wpm_unit']}"],
        [t["fillers"], str(metrics["filler_total"])],
    ]


def to_markdown(deck: Deck, metrics: dict, result: GradingResult, checks: list[FormalCheck], lang: str = "de") -> str:
    t = T.get(lang, T["de"])
    L = [f"# {t['title']}: {deck.source.stem}", "",
         f"**{t['total']}: {result.total_points:g} / {result.max_points:g} {t['points']}**", "",
         f"{t['recommendation']}: {result.recommendation_identified}", "", result.summary, "",
         f"## {t['criteria']}", "", f"| {t['criterion']} | {t['level']} | {t['points']} |", "|---|---|---|"]
    L += [f"| {c.name} | {t['levels'][c.level]} | {c.points:g} / {c.max_points:g} |" for c in result.criteria] + [""]
    for c in result.criteria:
        L += [f"### {c.name}: {t['levels'][c.level]}, {c.points:g} / {c.max_points:g}", "",
              f"**{t['holds']}** {c.was_traegt}", "", f"**{t['thin']}** {c.was_bleibt_duenn}", "",
              f"**{t['next']}** {c.naechster_schritt}", ""]
        if c.evidence:
            L += [f"{t['evidence']}:"] + [f"- {e}" for e in c.evidence] + [""]
    L += [f"## {t['strengths']}", ""] + [f"- {s}" for s in result.strengths] + ["", f"## {t['weaknesses']}", ""] + \
         [f"- {w}" for w in result.weaknesses] + [""]
    L += [f"## {t['coverage']}", "", f"| {t['task']} | {t['status']} | {t['note']} |", "|---|---|---|"]
    L += [f"| {c.label} | {t['statuses'][c.status]} | {c.note} |" for c in result.coverage]
    L += ["", f"## {t['delivery']}", ""] + [f"- **{d.label}.** {d.observation}" for d in result.delivery]
    L += ["", f"## {t['formal']}", ""] + [f"- {OK[c.ok]} {c.label}: {c.detail}" for c in checks]
    L += ["", f"## {t['metrics']}", ""] + [f"- {k}: {v}" for k, v in _metrics_rows(deck, metrics, t)]
    if result.slide_notes:
        L += ["", f"## {t['slide_notes']}", ""] + [f"- {t['slide']} {n.slide}: {n.observation}" for n in result.slide_notes]
    if result.flags:
        L += ["", f"## {t['flags']}", ""] + [f"- {f}" for f in result.flags]
    return "\n".join(L) + "\n"


def _table(doc, header: list[str], rows: list[list[str]]):
    tb = doc.add_table(rows=1, cols=len(header))
    tb.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        tb.rows[0].cells[i].text = h
    for r in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    return tb


def _bullets(doc, items: list[str]):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def to_docx(deck: Deck, metrics: dict, result: GradingResult, checks: list[FormalCheck], out: Path, lang: str = "de") -> Path:
    t = T.get(lang, T["de"])
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(f"{t['title']}: {deck.source.stem}", level=1)
    doc.add_paragraph().add_run(f"{t['total']}: {result.total_points:g} / {result.max_points:g} {t['points']}").bold = True
    p = doc.add_paragraph()
    p.add_run(f"{t['recommendation']}: ").italic = True
    p.add_run(result.recommendation_identified)
    doc.add_paragraph(result.summary)

    doc.add_heading(t["criteria"], level=2)
    _table(doc, [t["criterion"], t["level"], t["points"]],
           [[c.name, t["levels"][c.level], f"{c.points:g} / {c.max_points:g}"] for c in result.criteria])
    for c in result.criteria:
        doc.add_heading(f"{c.name}: {t['levels'][c.level]}, {c.points:g} / {c.max_points:g}", level=3)
        for label, text in ((t["holds"], c.was_traegt), (t["thin"], c.was_bleibt_duenn), (t["next"], c.naechster_schritt)):
            p = doc.add_paragraph()
            p.add_run(label + " ").bold = True
            p.add_run(text)
        if c.evidence:
            doc.add_paragraph(t["evidence"] + ":").runs[0].italic = True
            _bullets(doc, c.evidence)

    doc.add_heading(t["strengths"], level=2); _bullets(doc, result.strengths)
    doc.add_heading(t["weaknesses"], level=2); _bullets(doc, result.weaknesses)

    doc.add_heading(t["coverage"], level=2)
    _table(doc, [t["task"], t["status"], t["note"]], [[c.label, t["statuses"][c.status], c.note] for c in result.coverage])

    doc.add_heading(t["delivery"], level=2)
    for d in result.delivery:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(d.label + ". ").bold = True
        p.add_run(d.observation)

    doc.add_heading(t["formal"], level=2)
    _bullets(doc, [f"{OK[c.ok]} {c.label}: {c.detail}" for c in checks])

    doc.add_heading(t["metrics"], level=2)
    _table(doc, [t["metric"], t["value"]], _metrics_rows(deck, metrics, t))

    if result.slide_notes:
        doc.add_heading(t["slide_notes"], level=2)
        _bullets(doc, [f"{t['slide']} {n.slide}: {n.observation}" for n in result.slide_notes])
    if result.flags:
        doc.add_heading(t["flags"], level=2)
        _bullets(doc, result.flags)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
