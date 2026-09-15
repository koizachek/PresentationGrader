"""Grading context (task, rubric, case) with per-job overrides.

Defaults come from the files in app/rubrics (or RUBRIC_PATH / TASK_PATH /
CASE_PATH). A job may upload replacements; text is extracted from md, txt,
pdf, pptx or docx, rubrics from yaml or json.
"""
from __future__ import annotations

import json
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import yaml

from .config import settings

TEXT_EXT = {".md", ".txt", ".pdf", ".pptx", ".docx"}
RUBRIC_EXT = {".yaml", ".yml", ".json"}


@dataclass
class GradingContext:
    rubric: dict
    task: str
    case: str
    rubric_source: str
    task_source: str
    case_source: str

    def summary(self) -> dict:
        return {
            "rubric": {"title": self.rubric.get("meta", {}).get("course") or self.rubric.get("title", ""),
                       "version": self.rubric.get("meta", {}).get("version", ""),
                       "max_points": self.rubric.get("max_points"),
                       "criteria": [c.get("name") for c in self.rubric.get("criteria", [])],
                       "source": self.rubric_source},
            "task": {"words": len(self.task.split()), "first_line": self.task.strip().splitlines()[0] if self.task.strip() else "",
                     "source": self.task_source},
            "case": {"words": len(self.case.split()), "first_line": self.case.strip().splitlines()[0] if self.case.strip() else "",
                     "source": self.case_source},
        }


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        if shutil.which("pdftotext"):
            r = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, timeout=120)
            if r.returncode == 0 and r.stdout.strip():
                return r.stdout
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            lines = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            for sh in slide.shapes:
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))
            if lines:
                out.append(f"## Folie {i}\n" + "\n".join(lines))
        return "\n\n".join(out)
    if ext == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {ext} (allowed: {', '.join(sorted(TEXT_EXT))})")


def load_rubric_file(path: Path) -> dict:
    ext = path.suffix.lower()
    raw = path.read_text(encoding="utf-8")
    data = json.loads(raw) if ext == ".json" else yaml.safe_load(raw)
    validate_rubric(data)
    return data


def validate_rubric(data) -> None:
    if not isinstance(data, dict):
        raise ValueError("Rubrik muss ein Objekt sein (YAML/JSON mit 'criteria' und 'max_points').")
    crit = data.get("criteria")
    if not isinstance(crit, list) or not crit:
        raise ValueError("Rubrik braucht eine nicht-leere Liste 'criteria'.")
    total = 0.0
    for i, c in enumerate(crit, 1):
        for key in ("id", "name", "max_points"):
            if key not in c:
                raise ValueError(f"Kriterium {i}: Feld '{key}' fehlt.")
        levels = c.get("levels")
        if not isinstance(levels, dict) or not {"ueberzeugend", "tragfaehig", "ansatzweise"} <= set(levels):
            raise ValueError(f"Kriterium '{c['id']}': 'levels' braucht ueberzeugend, tragfaehig, ansatzweise.")
        total += float(c["max_points"])
    if "max_points" in data and abs(float(data["max_points"]) - total) > 1e-6:
        raise ValueError(f"max_points ({data['max_points']}) entspricht nicht der Summe der Kriterien ({total:g}).")
    for key in ("coverage", "delivery_aspects"):
        if key in data and not isinstance(data[key], list):
            raise ValueError(f"'{key}' muss eine Liste sein.")


def _default_rubric() -> dict:
    return yaml.safe_load(settings.rubric_path.read_text(encoding="utf-8"))


def _read(p: Path) -> str:
    return p.read_text(encoding="utf-8") if p.exists() else ""


def default_context() -> GradingContext:
    return GradingContext(
        rubric=_default_rubric(), task=_read(settings.task_path), case=_read(settings.case_path),
        rubric_source=settings.rubric_path.name, task_source=settings.task_path.name, case_source=settings.case_path.name,
    )


def build_context(job_dir: Path, rubric: Path | None = None, task: Path | None = None, case: Path | None = None) -> GradingContext:
    ctx = default_context()
    if rubric:
        ctx.rubric, ctx.rubric_source = load_rubric_file(rubric), rubric.name
    if task:
        ctx.task, ctx.task_source = extract_text(task), task.name
    if case:
        ctx.case, ctx.case_source = extract_text(case), case.name
    return ctx
